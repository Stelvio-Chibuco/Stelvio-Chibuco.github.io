"""Gera os CVs (PT e EN) em .docx com o mesmo aspecto dos PDFs publicados.

Uso:
    pip install python-docx
    python3 cv/gerar_cv.py
    soffice --headless --convert-to pdf --outdir public/cv cv/*.docx

Para mudar o conteúdo, edite os dicionários PT e EN abaixo (ou abra o .docx
gerado no Word/LibreOffice) e exporte de novo para PDF.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x4E, 0x79)
CINZENTO = RGBColor(0x55, 0x55, 0x55)
QUASE_PRETO = RGBColor(0x1A, 0x1A, 0x1A)
FONTE = "Liberation Sans"
LARGURA_TEXTO = Cm(21.0 - 2 * 2.0)

LINKS = {
    "email": "mailto:stelviochibuco@outlook.com",
    "portfolio": "https://stelvio-chibuco.netlify.app/",
    "github": "https://github.com/Stelvio-Chibuco",
    "linkedin": "https://www.linkedin.com/in/st%C3%A9lvio-chibuco-301676263",
    "boleia": "https://admin.16-16-231-39.sslip.io/login",
}


def hyperlink(paragraph, text, url, size=10):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F4E79"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        props.append(el)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def borda_inferior(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    borders.append(bottom)
    p_pr.append(borders)


def espacamento_letras(run, pontos):
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(int(pontos * 20)))
    run._r.get_or_add_rPr().append(spacing)


def paragrafo(doc, antes=0, depois=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)
    return p


def texto(p, conteudo, bold=False, italic=False, size=10, color=None):
    run = p.add_run(conteudo)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def rico(p, partes, size=10):
    """partes: lista de str ou ('link', texto, chave)."""
    for parte in partes:
        if isinstance(parte, tuple):
            hyperlink(p, parte[1], LINKS[parte[2]], size)
        else:
            texto(p, parte, size=size)


def titulo_seccao(doc, nome):
    p = paragrafo(doc, antes=12, depois=6)
    p.paragraph_format.keep_with_next = True
    run = texto(p, nome.upper(), bold=True, size=11, color=AZUL)
    espacamento_letras(run, 1)
    borda_inferior(p)


def entrada(doc, titulo, direita, antes=6):
    p = paragrafo(doc, antes=antes)
    p.paragraph_format.tab_stops.add_tab_stop(LARGURA_TEXTO, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    texto(p, titulo, bold=True)
    if direita:
        texto(p, "\t" + direita, color=CINZENTO)


def linha_italica(doc, conteudo):
    p = paragrafo(doc, depois=1)
    p.paragraph_format.keep_with_next = True
    texto(p, conteudo, italic=True, size=9.5, color=CINZENTO)


def marcador(doc, partes):
    p = paragrafo(doc, depois=1)
    fmt = p.paragraph_format
    fmt.left_indent = Pt(17)
    fmt.first_line_indent = Pt(-11)
    fmt.tab_stops.add_tab_stop(Pt(17))
    texto(p, "•\t")
    rico(p, partes if isinstance(partes, list) else [partes])


def linha_rotulo(doc, rotulo, conteudo):
    p = paragrafo(doc, depois=2)
    texto(p, rotulo + ": ", bold=True)
    texto(p, conteudo)


def gerar(d, destino):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    props = doc.core_properties
    props.title = d["doc_title"]
    props.author = "Stélvio André Chibuco"

    p = paragrafo(doc, depois=4)
    texto(p, "Stélvio André Chibuco", bold=True, size=20, color=QUASE_PRETO)
    p = paragrafo(doc, depois=4)
    texto(p, d["headline"], bold=True, size=11, color=AZUL)
    for linha in d["contact"]:
        p = paragrafo(doc, depois=1)
        rico(p, linha)

    titulo_seccao(doc, d["h_profile"])
    p = paragrafo(doc)
    texto(p, d["profile"])

    titulo_seccao(doc, d["h_projects"])
    for i, proj in enumerate(d["projects"]):
        entrada(doc, proj["title"], proj["kind"], antes=0 if i == 0 else 6)
        linha_italica(doc, proj["stack"])
        for b in proj["bullets"]:
            marcador(doc, b)

    titulo_seccao(doc, d["h_experience"])
    for i, exp in enumerate(d["experience"]):
        entrada(doc, exp["title"], exp["period"], antes=0 if i == 0 else 6)
        for b in exp["bullets"]:
            marcador(doc, b)

    titulo_seccao(doc, d["h_education"])
    entrada(doc, d["degree"], d["degree_period"], antes=0)
    linha_italica(doc, d["school"])
    for b in d["education_bullets"]:
        marcador(doc, b)

    titulo_seccao(doc, d["h_courses"])
    for rotulo, conteudo in d["courses"]:
        linha_rotulo(doc, rotulo, conteudo)

    titulo_seccao(doc, d["h_skills"])
    for rotulo, conteudo in d["skills"]:
        linha_rotulo(doc, rotulo, conteudo)

    titulo_seccao(doc, d["h_other"])
    for rotulo, conteudo in d["other"]:
        linha_rotulo(doc, rotulo, conteudo)

    doc.save(destino)


SEP = "  ·  "

PT = {
    "doc_title": "Curriculum Vitae — Stélvio André Chibuco",
    "headline": "Desenvolvedor Full-Stack · Python/Django, React, Flutter · Visão computacional",
    "contact": [
        ["Xai-Xai, Moçambique" + SEP + "+258 84 766 9125 (WhatsApp)" + SEP + "+258 82 766 9125"],
        ["E-mail: ", ("link", "stelviochibuco@outlook.com", "email")],
        [
            "Portfólio: ",
            ("link", "stelvio-chibuco.netlify.app", "portfolio"),
            SEP + "GitHub: ",
            ("link", "github.com/Stelvio-Chibuco", "github"),
            SEP,
            ("link", "LinkedIn", "linkedin"),
        ],
    ],
    "h_profile": "Perfil",
    "profile": (
        "Desenvolvedor de sistemas, licenciado em Informática (Engenharia de Desenvolvimento de "
        "Sistemas, Universidade Save, 2024), que constrói produtos web e mobile completos, do backend "
        "à publicação. Desenvolvi uma loja online com pagamentos M-Pesa e e-Mola, uma plataforma de "
        "transporte por app com rastreamento em tempo real e um sistema SaaS de controlo de acesso por "
        "reconhecimento facial. Tenho experiência em suporte técnico, redes e servidores Linux, e "
        "formação em cibersegurança pela Cisco Networking Academy."
    ),
    "h_projects": "Projectos",
    "projects": [
        {
            "title": "SaveGuard — Controlo de acesso por reconhecimento facial",
            "kind": "Produto SaaS",
            "stack": "FastAPI · React · TypeScript · TensorFlow.js · PostgreSQL · WebSockets · Docker",
            "bullets": [
                "Reconhecimento facial no browser com prova de vida (piscar de olhos) contra fotografias; "
                "identidade verificada no servidor antes de abrir a porta.",
                "Integração com fechaduras reais (relés Wi-Fi Shelly, relés USB, Raspberry Pi) e alertas "
                "com fotografia por e-mail, Telegram e webhook.",
                "SaaS multi-cliente: equipas e papéis, planos com teste grátis, modo quiosque e painel de "
                "segurança em tempo real.",
            ],
        },
        {
            "title": "Motice Congelados & Diversos — Loja online",
            "kind": "Projecto para cliente",
            "stack": "Django · Django REST · PostgreSQL · Redis · PWA · Docker",
            "bullets": [
                "Loja mobile-first instalável como app (PWA), rápida em redes lentas.",
                "Pagamentos M-Pesa (API Vodacom C2B), e-Mola, mKesh, PayPal/cartão e pagamento na "
                "entrega, todos verificados no servidor.",
                "Zonas de entrega com taxas, localização do cliente no mapa e painel de gestão pensado "
                "para o telemóvel.",
            ],
        },
        {
            "title": "Boleia+ · Chibuto Ride — Plataforma de mobilidade",
            "kind": "Startup · testes fechados",
            "stack": "Django · DRF · PostGIS · Celery · Redis · Flutter · Next.js · Leaflet",
            "bullets": [
                "App Flutter para passageiro e motorista: registo por telefone com OTP, viagem em tempo "
                "real, PIN de embarque e botão SOS.",
                [
                    "API geoespacial e dashboard de operações com mapas em tempo real (",
                    ("link", "em produção", "boleia"),
                    ").",
                ],
            ],
        },
        {
            "title": "Stock & Venda — Sistema de stock e vendas",
            "kind": "Sistema de gestão",
            "stack": "PHP · MySQL · JavaScript · Bootstrap",
            "bullets": [
                "Catálogo, entradas de stock e vendas com actualização automática, contas de utilizador "
                "e relatórios de vendas.",
            ],
        },
    ],
    "h_experience": "Experiência profissional",
    "experience": [
        {
            "title": "Desenvolvedor Full-Stack Freelancer — Independente",
            "period": "2023 — presente",
            "bullets": [
                "Desenvolvimento de sistemas de e-commerce, mobilidade e controlo de acesso para "
                "clientes e produtos próprios."
            ],
        },
        {
            "title": "Estagiário Técnico-Profissional — INAGE, Instituto Nacional de Governo Electrónico (Gaza)",
            "period": "2025",
            "bullets": ["Formador substituto do curso de Informática Básica e suporte técnico informático."],
        },
        {
            "title": "Estagiário Técnico-Profissional — Conselho Municipal da Cidade de Xai-Xai",
            "period": "2024 — 2025",
            "bullets": ["Suporte técnico aos serviços municipais: equipamento, software, redes e utilizadores."],
        },
        {
            "title": "Estagiário Técnico-Profissional (Desenvolvimento de Sistemas) — Universidade Save",
            "period": "2023",
            "bullets": ["Instalação e configuração do servidor de e-mail da universidade (Zimbra)."],
        },
    ],
    "h_education": "Formação académica",
    "degree": "Licenciatura em Informática — Engenharia de Desenvolvimento de Sistemas",
    "degree_period": "2020 — 2024",
    "school": "Universidade Save",
    "education_bullets": [
        "Projecto final: controlo de acesso por reconhecimento facial com alertas de segurança.",
        "Apresentação do projecto nas II Jornadas Científicas da Universidade Joaquim Chissano (2023).",
    ],
    "h_courses": "Cursos — Cisco Networking Academy",
    "courses": [
        (
            "2025",
            "Hacker Ético · Gestão de Ameaças Cibernéticas · Fundamentos de Redes · Fundamentos do "
            "Python I · Inglês para TI II",
        ),
        ("2023", "Introdução à Cibersegurança · Noções Básicas de Redes"),
    ],
    "h_skills": "Competências técnicas",
    "skills": [
        ("Backend", "Python, Django, Django REST, FastAPI, Java, PHP, C#"),
        ("Frontend e mobile", "React, Next.js, TypeScript, Tailwind CSS, Bootstrap, PWA, Flutter/Dart"),
        ("Bases de dados", "PostgreSQL, PostGIS, MySQL, Redis, SQLite"),
        (
            "DevOps e sistemas",
            "Docker, Nginx, Git, CI/CD, Linux (Debian, Ubuntu, Kali), Windows, servidor de e-mail Zimbra",
        ),
        ("IA e segurança", "Visão computacional, TensorFlow.js, OpenCV, hacking ético, redes de computadores"),
    ],
    "h_other": "Línguas e outras informações",
    "other": [
        ("Línguas", "Português (nativo), Changana (nativo), Inglês (profissional)"),
        ("Blog", "artigos e tutoriais sobre cibersegurança e Linux em cyber-gorilla.blogspot.com"),
        (
            "Outros",
            "Carta de condução de ligeiros e pesados · Serviço militar cumprido (2015 — 2016) · "
            "Disponibilidade imediata",
        ),
    ],
}

EN = {
    "doc_title": "Résumé — Stélvio André Chibuco",
    "headline": "Full-Stack Software Developer · Python/Django, React, Flutter · Computer vision",
    "contact": [
        ["Xai-Xai, Mozambique (UTC+2)" + SEP + "+258 84 766 9125 (WhatsApp)" + SEP + "+258 82 766 9125"],
        ["Email: ", ("link", "stelviochibuco@outlook.com", "email"), SEP + "Open to remote work"],
        [
            "Portfolio: ",
            ("link", "stelvio-chibuco.netlify.app", "portfolio"),
            SEP + "GitHub: ",
            ("link", "github.com/Stelvio-Chibuco", "github"),
            SEP,
            ("link", "LinkedIn", "linkedin"),
        ],
    ],
    "h_profile": "Profile",
    "profile": (
        "Systems developer with a BSc in Informatics (Systems Development Engineering, Universidade "
        "Save, 2024) who builds complete web and mobile products, from the backend to deployment. I have "
        "built an online store with M-Pesa and e-Mola mobile-money payments, a ride-hailing platform "
        "with real-time tracking and a multi-tenant SaaS for facial-recognition access control. "
        "Background in IT support, networking and Linux servers, with cybersecurity training from Cisco "
        "Networking Academy."
    ),
    "h_projects": "Projects",
    "projects": [
        {
            "title": "SaveGuard — Facial-recognition access control",
            "kind": "SaaS product",
            "stack": "FastAPI · React · TypeScript · TensorFlow.js · PostgreSQL · WebSockets · Docker",
            "bullets": [
                "In-browser face recognition with a blink liveness check against photo spoofing; identity "
                "re-verified on the server before the door opens.",
                "Drives real locks (Shelly Wi-Fi relays, USB relays, Raspberry Pi) and sends photo alerts "
                "by email, Telegram and webhook.",
                "Multi-tenant SaaS: teams and roles, plans with a free trial, door kiosk mode and a "
                "real-time security dashboard.",
            ],
        },
        {
            "title": "Motice Congelados & Diversos — Online store",
            "kind": "Client project",
            "stack": "Django · Django REST · PostgreSQL · Redis · PWA · Docker",
            "bullets": [
                "Mobile-first store installable as an app (PWA), fast on slow networks.",
                "Payments via M-Pesa (Vodacom C2B API), e-Mola, mKesh, PayPal/card and cash on delivery, "
                "all verified server-side.",
                "Delivery zones with fees, customer location on a map and a phone-first admin panel.",
            ],
        },
        {
            "title": "Boleia+ · Chibuto Ride — Mobility platform",
            "kind": "Startup · closed beta",
            "stack": "Django · DRF · PostGIS · Celery · Redis · Flutter · Next.js · Leaflet",
            "bullets": [
                "Flutter passenger and driver app: phone + OTP sign-up, live trip tracking, boarding PIN "
                "and an SOS button.",
                [
                    "Geospatial API and an operations dashboard with live maps (",
                    ("link", "in production", "boleia"),
                    ").",
                ],
            ],
        },
        {
            "title": "Stock & Venda — Inventory and point of sale",
            "kind": "Business system",
            "stack": "PHP · MySQL · JavaScript · Bootstrap",
            "bullets": [
                "Product catalogue, stock intake and sales with automatic stock updates, user accounts "
                "and sales reports.",
            ],
        },
    ],
    "h_experience": "Experience",
    "experience": [
        {
            "title": "Freelance Full-Stack Developer — Independent",
            "period": "2023 — present",
            "bullets": ["E-commerce, mobility and access-control systems for clients and my own products."],
        },
        {
            "title": "Technical-Professional Intern — INAGE, National Institute of e-Government (Gaza)",
            "period": "2025",
            "bullets": ["Substitute trainer for the Basic Computing course and IT technical support."],
        },
        {
            "title": "Technical-Professional Intern — Xai-Xai City Council",
            "period": "2024 — 2025",
            "bullets": ["IT support for municipal services: hardware, software, networks and users."],
        },
        {
            "title": "Technical-Professional Intern (Systems Development) — Universidade Save",
            "period": "2023",
            "bullets": ["Installed and configured the university e-mail server (Zimbra)."],
        },
    ],
    "h_education": "Education",
    "degree": "BSc in Informatics — Systems Development Engineering",
    "degree_period": "2020 — 2024",
    "school": "Universidade Save, Mozambique",
    "education_bullets": [
        "Final project: facial-recognition access control with security alerts.",
        "Presented the project at the II Scientific Conference of Universidade Joaquim Chissano (2023).",
    ],
    "h_courses": "Courses — Cisco Networking Academy",
    "courses": [
        (
            "2025",
            "Ethical Hacker · Cyber Threat Management · Networking Essentials · Python Essentials 1 · "
            "English for IT 2",
        ),
        ("2023", "Introduction to Cybersecurity · Networking Basics"),
    ],
    "h_skills": "Technical skills",
    "skills": [
        ("Backend", "Python, Django, Django REST, FastAPI, Java, PHP, C#"),
        ("Frontend & mobile", "React, Next.js, TypeScript, Tailwind CSS, Bootstrap, PWA, Flutter/Dart"),
        ("Databases", "PostgreSQL, PostGIS, MySQL, Redis, SQLite"),
        ("DevOps & systems", "Docker, Nginx, Git, CI/CD, Linux (Debian, Ubuntu, Kali), Windows, Zimbra mail server"),
        ("AI & security", "Computer vision, TensorFlow.js, OpenCV, ethical hacking, computer networks"),
    ],
    "h_other": "Languages & other",
    "other": [
        ("Languages", "Portuguese (native), Changana (native), English (professional)"),
        ("Blog", "cybersecurity and Linux articles and tutorials at cyber-gorilla.blogspot.com"),
        (
            "Other",
            "Driving licence (light and heavy vehicles) · Military service completed (2015 — 2016) · "
            "Available immediately",
        ),
    ],
}


if __name__ == "__main__":
    pasta = Path(__file__).resolve().parent
    gerar(PT, pasta / "Stelvio-Chibuco-CV-PT.docx")
    gerar(EN, pasta / "Stelvio-Chibuco-CV-EN.docx")
    print("Gerados:", *(p.name for p in sorted(pasta.glob("*.docx"))))
